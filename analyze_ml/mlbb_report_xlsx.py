# -*- coding: utf-8 -*-
"""Build mlbb_swiss_report.xlsx from the analysis CSVs.

Proper Excel workbook (openpyxl): score-like columns (Series, Games, W/L,
K/D/A) are written as TEXT cells so Excel can never turn "3-0" into a date.
Re-run after any pipeline change.
"""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE = r"D:/Desktop/OBS Mobile Legends ML Setup/analyze_ml"
OUT = BASE + "/mlbb_swiss_report.xlsx"

SHEETS = [
    ("Team Standings", "mlbb_team_standings.csv"),
    ("Player MVP", "mlbb_player_mvp_standings.csv"),
    ("Most Played Heroes", "mlbb_most_played_heroes.csv"),
]

# headers that must stay TEXT even if they look numeric/date-like
TEXT_HEADERS = {"Series", "Games", "W/L", "K/D/A", "Series%", "Game%", "Win%", "TopHeroID"}

HDR_FILL = PatternFill("solid", fgColor="1F4E5F")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)

def clean(v):
    """Strip the CSV ="..." Excel-text wrapper back to the plain value."""
    if isinstance(v, str) and v.startswith('="') and v.endswith('"'):
        return v[2:-1]
    return v

def num(v):
    """Best-effort numeric conversion; falls back to the raw value."""
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return v

wb = Workbook()
wb.remove(wb.active)
for title, fname in SHEETS:
    rows = list(csv.DictReader(open(f"{BASE}/{fname}", encoding="utf-8-sig")))
    if not rows:
        continue
    ws = wb.create_sheet(title)
    headers = list(rows[0].keys())
    ws.append(headers)
    for r in rows:
        vals = []
        for h in headers:
            v = clean(r[h])
            vals.append(v if h in TEXT_HEADERS else num(v))
        ws.append(vals)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HDR_FILL, HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for c, h in enumerate(headers, 1):
        if h in TEXT_HEADERS:
            for row in ws.iter_rows(min_row=2, min_col=c, max_col=c):
                for cell in row:
                    cell.number_format = "@"
        width = max([len(str(headers[c - 1]))] +
                    [len(str(ws.cell(row=i, column=c).value or "")) for i in range(2, ws.max_row + 1)])
        ws.column_dimensions[get_column_letter(c)].width = min(width + 2, 30)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

wb.save(OUT)
print(f"Saved: {OUT}  sheets={[s for s, _ in SHEETS]}")
